"""Extract game-specific data from the Run & Bun master sheet into data/game_data.json.

Game data = things that are the same for every attempt: the route list and each
route's encounter table (Encounters tab), every enemy trainer and their team
(Trainer Lookup tab), the level-cap fights (header rows of the Past Boss Fights
tab), and the dupes-clause evolution families (helper table at the bottom of
the Encounters tab).

Usage:
    python scripts/extract_game_data.py ["path/to/sheet.xlsx"]
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "Copy of Run & Bun Master Sheet.xlsx"
OUT = ROOT / "data" / "game_data.json"

# --- Encounters tab layout -------------------------------------------------
HEADER_ROW = 2
SECTIONS = [  # (method, first row, last row)
    ("Land", 5, 16),
    ("Fishing", 17, 26),
    ("Surf", 27, 31),
    ("Rocks", 32, 36),
    ("Other", 37, 58),
]
STARTER_COL = 2       # Starter block: Level, Pokemon, checkbox
FIRST_ROUTE_COL = 5   # Every route block after it: %, Level, Pokemon, checkbox
ROUTE_BLOCK_WIDTH = 4
FAMILY_ROWS = range(94, 131)  # one evolution family per column
FAMILY_FIRST_COL = 3

# --- Past Boss Fights tab layout -------------------------------------------
BOSS_LEVEL_ROW = 1
BOSS_NAME_ROW = 2
BOSS_FIRST_COL = 2

# --- Trainer Lookup tab layout ---------------------------------------------
# One trainer per row; each Pokemon cell is
# "Species,Level,Item,Ability,Nature,Move1,Move2,Move3,Move4".
TRAINER_FIRST_ROW = 2
TRAINER_LOCATION_COL = 3
TRAINER_NAME_COL = 4
TRAINER_POKEMON_COLS = range(6, 12)

# Level-cap fights (Past Boss Fights names) -> the Trainer Lookup rows fought
# there, as (trainer name, location). Several rows means back-to-back or
# double battles (Museum grunts, Tate & Liza) or alternative teams (the rival's
# team depends on your starter).
LEVEL_CAP_TRAINERS = {
    "Route 104 Aqua Grunt": [("Team Aqua Grunt Petalburg Woods [Boss]", "Route 104")],
    "Museum Aqua Grunts": [("Team Aqua Grunt Museum #1 [Boss]", "Slateport Museum"),
                           ("Team Aqua Grunt Museum #2 [Boss]", "Slateport Museum")],
    "Leader Brawly": [("Leader Brawly [Boss]", "Dewford Gym")],
    "Leader Roxanne": [("Leader Roxanne [Boss]", "Rustboro Gym")],
    "Route 117 Chelle": [("Trainer Chelle Daycare [Boss]", "Route 117")],
    "Leader Wattson": [("Leader Wattson [Boss]", "Mauville Gym")],
    "Cycling Road Rival": [(f"Trainer Rival Cycling Road {s} [Boss]", "Route 110")
                           for s in ("Sceptile", "Blaziken", "Swampert")],
    "Leader Norman": [("Leader Norman [Boss]", "Petalburg Gym")],
    "Fallarbor Town Vito": [("Winstrate Vito [Boss]", "Fallarbor")],
    "Mt. Chimney Maxie": [("Magma Leader Maxie [Boss]", "Mt. Chimney")],
    "Leader Flannery": [("Leader Flannery [Double] [Boss]", "Lavaridge Gym")],
    "Weather Institute Shelly": [("Aqua Admin Shelly Weather Institute [Boss]", "Weather Institute")],
    "Route 119 Rival": [(f"Trainer Rival Bridge {s} [Double] [Boss]", "Route 119")
                        for s in ("Sceptile", "Blaziken", "Swampert")],
    "Leader Winona": [("Leader Winona [Boss]", "Fortree Gym")],
    "Lilycove City Rival": [(f"Trainer Rival Lilycove {s} [Boss]", "Lilycove")
                            for s in ("Sceptile", "Blaziken", "Swampert")],
    "Mt. Pyre Archie": [("Aqua Leader Archie [Tag Battle] [Boss]", "Mt. Pyre")],
    "Magma Hideout Maxie": [("Magma Leader Maxie [Boss]", "Magma Hideout")],
    "Aqua Hideout Matt": [("Aqua Admin Matt [Boss]", "Aqua Hideout")],
    "Leaders Tate & Liza": [("Leader Tate [Boss]", "Mossdeep Gym"),
                            ("Leader Liza [Boss]", "Mossdeep Gym")],
    "Seafloor Cavern Archie": [("Aqua Leader Archie [Boss]", "Seafloor Cavern")],
    "Leader Juan": [("Leader Juan [Double] [Boss]", "Sootopolis Gym")],
    "Victory Road Vito": [("Winstrate Vito [Boss]", "Victory Road")],
    "Champion Wallace": [("Champion Wallace", "Pokémon League")],
}

# Level-cap fights against two trainers with separate teams (rather than
# alternative teams). Each trainer becomes a battle of its own, so KOs are
# tracked against one team at a time. The last one is the split's level-cap
# battle; the others are listed just before it in the same split.
SEPARATE_TRAINERS = {"Museum Aqua Grunts", "Leaders Tate & Liza"}


def clean(value):
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def normalize_rate(value):
    """'20%' -> 20.0, 0.0909 -> 9.09 (fractions are used for gift/fossil odds)."""
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, str):
        return float(value.rstrip("%"))
    if isinstance(value, (int, float)):
        return round(value * 100, 2) if value <= 1 else float(value)
    return None


def normalize_level(value):
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_block(ws, rate_col, level_col, pokemon_col, methods=None):
    encounters = []
    for method, first, last in SECTIONS:
        if methods and method not in methods:
            continue
        for row in range(first, last + 1):
            pokemon = clean(ws.cell(row, pokemon_col).value)
            if not isinstance(pokemon, str):
                continue
            encounters.append({
                "method": method,
                "pokemon": pokemon,
                "rate": normalize_rate(ws.cell(row, rate_col).value) if rate_col else None,
                "level": normalize_level(ws.cell(row, level_col).value),
            })
    return encounters


def extract_routes(ws):
    # The Starter column's upper rows hold rule-settings questions, not Pokemon.
    routes = [{
        "name": clean(ws.cell(HEADER_ROW, STARTER_COL).value),
        "encounters": read_block(ws, None, STARTER_COL, STARTER_COL + 1, methods={"Other"}),
    }]
    col = FIRST_ROUTE_COL
    while (name := clean(ws.cell(HEADER_ROW, col).value)) is not None:
        routes.append({"name": name, "encounters": read_block(ws, col, col + 1, col + 2)})
        col += ROUTE_BLOCK_WIDTH
    return routes


def extract_families(ws):
    families = []
    for col in range(FAMILY_FIRST_COL, ws.max_column + 1):
        members = [clean(ws.cell(row, col).value) for row in FAMILY_ROWS]
        members = [m for m in members if isinstance(m, str)]
        if members:
            families.append(members)
    return families


def extract_bosses(ws):
    bosses = []
    col = BOSS_FIRST_COL
    while (name := clean(ws.cell(BOSS_NAME_ROW, col).value)) is not None:
        level = clean(ws.cell(BOSS_LEVEL_ROW, col).value)
        level_cap = int(level.lower().replace("lvl", "").strip()) if isinstance(level, str) else None
        bosses.append({"name": name, "level_cap": level_cap})
        col += 1
    return bosses


def parse_trainer_pokemon(value):
    """'Kubfu,20 ,Iapapa Berry,Inner Focus,Jolly,Brick Break,...' -> dict."""
    species, level, item, ability, nature, *moves = [part.strip() for part in value.split(",")]
    return {
        "species": species,
        "level": int(level) if level.isdigit() else None,
        "item": item or None,
        "ability": ability or None,
        "nature": nature or None,
        "moves": [m for m in moves if m],
    }


def extract_trainers(ws):
    """Every trainer on the Trainer Lookup tab, in game order."""
    trainers = []
    for row in range(TRAINER_FIRST_ROW, ws.max_row + 1):
        raw_name = clean(ws.cell(row, TRAINER_NAME_COL).value)
        if raw_name is None:
            continue
        pokemon = [clean(ws.cell(row, col).value) for col in TRAINER_POKEMON_COLS]
        trainers.append({
            "name": re.sub(r"\s*\[[^\]]*\]", "", raw_name).strip(),
            "raw_name": raw_name,
            "location": clean(ws.cell(row, TRAINER_LOCATION_COL).value),
            "key": f"{raw_name} @ {clean(ws.cell(row, TRAINER_LOCATION_COL).value)}",
            "tags": re.findall(r"\[([^\]]+)\]", raw_name),
            "pokemon": [parse_trainer_pokemon(p) for p in pokemon if p],
        })
    # Names can repeat once tags are stripped (Elite Four singles vs "[Double]"
    # teams), so trainers are identified by their full sheet name + location.
    keys = Counter(t["key"] for t in trainers)
    duplicates = [k for k, n in keys.items() if n > 1]
    if duplicates:
        raise SystemExit(f"Trainer name + location must be unique: {duplicates}")
    return trainers


def build_battles(bosses, trainers):
    """Group trainers into battles, in game order.

    The level-cap fights use the trainers listed in LEVEL_CAP_TRAINERS, one
    battle per trainer for SEPARATE_TRAINERS. Every other trainer is its own
    battle, listed under the next level-cap fight. A "[Tag Partner]" (an ally)
    joins the battle of the trainer just before it.

    `split` names the split a level-cap battle ends (its Past Boss Fights
    name); `group` names the split every other battle is listed in.
    """
    by_key = {(t["raw_name"], t["location"]): t for t in trainers}
    boss_of = {}
    for boss in bosses:
        for key in LEVEL_CAP_TRAINERS[boss["name"]]:
            if key not in by_key:
                raise SystemExit(f"{boss['name']}: no trainer {key} on the Trainer Lookup tab")
            boss_of[key] = boss

    battles, boss_battles, waiting = [], {}, []
    for trainer in trainers:
        boss = boss_of.get((trainer["raw_name"], trainer["location"]))
        if "Tag Partner" in trainer["tags"] and battles:
            battles[-1]["trainers"].append(trainer)
            continue
        if boss is None:
            battle = {"name": trainer["name"], "location": trainer["location"], "level_cap": None,
                      "group": None, "split": None, "trainers": [trainer]}
            battles.append(battle)
            waiting.append(battle)
            continue
        for other in waiting:
            other["group"] = boss["name"]
        waiting = []
        separate = boss["name"] in SEPARATE_TRAINERS
        if boss["name"] in boss_battles and not separate:
            boss_battles[boss["name"]][-1]["trainers"].append(trainer)  # an alternative team
            continue
        battle = {"name": trainer["name"] if separate else boss["name"], "location": trainer["location"],
                  "level_cap": boss["level_cap"], "group": None, "split": boss["name"], "trainers": [trainer]}
        boss_battles.setdefault(boss["name"], []).append(battle)
        battles.append(battle)
    if waiting:
        raise SystemExit(f"Trainers after the last level-cap fight: {[b['name'] for b in waiting]}")
    for name, parts in boss_battles.items():
        for part in parts[:-1]:  # the last part is the split's level-cap battle
            part["group"], part["split"] = name, None
    for battle in battles:
        # A split's level-cap battle keeps its Past Boss Fights name as its identity.
        battle["key"] = (f"boss:{battle['split']}" if battle["split"] is not None
                         else f"trainer:{battle['trainers'][0]['key']}")
        for trainer in battle["trainers"]:
            del trainer["raw_name"]
    return battles


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    bosses = extract_bosses(wb["Past Boss Fights"])
    trainers = extract_trainers(wb["Trainer Lookup"])
    data = {
        "source": xlsx.name,
        "routes": extract_routes(wb["Encounters"]),
        "battles": build_battles(bosses, trainers),
        "families": extract_families(wb["Encounters"]),
    }
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    n_enc = sum(len(r["encounters"]) for r in data["routes"])
    n_boss = sum(1 for b in data["battles"] if b["split"] is not None)
    print(f"Wrote {OUT}: {len(data['routes'])} routes, {n_enc} encounter slots, "
          f"{len(data['battles'])} battles ({n_boss} splits) from {len(trainers)} trainers, "
          f"{len(data['families'])} families")


if __name__ == "__main__":
    main()
